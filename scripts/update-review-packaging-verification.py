"""Update the final packaging verification after the WiX warning fixes."""

import io
import os
from pathlib import Path
import tempfile
import zipfile

from docx import Document
from lxml import etree


report = Path(__file__).resolve().parents[1] / "项目审查报告-2026-09-04.docx"
document = Document(report)
prefix = "维护：签名流程故障注入通过；"
matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
if len(matches) != 1 or len(matches[0].runs) != 1:
    raise ValueError("Expected one single-run packaging verification paragraph")
matches[0].runs[0].text = (
    "维护：签名流程故障注入通过；WiX 临时不可安装夹具完成编译、ICE 和安装表检查，"
    "未执行产品安装。构建无 WiX 或 ICE 警告，四个未版本化 EXE 进入 MsiFileHash，"
    "快捷方式不跨组件引用文件。"
)
for field in (
    "title", "subject", "author", "last_modified_by", "comments", "keywords",
    "category", "identifier", "language", "content_status", "version",
):
    setattr(document.core_properties, field, "")

buffer = io.BytesIO()
document.save(buffer)
temporary = None
try:
    with tempfile.NamedTemporaryFile(
        dir=report.parent, prefix=".review-packaging-", suffix=".docx", delete=False,
    ) as output:
        temporary = Path(output.name)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root:
                    if etree.QName(node).localname in (
                        "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
                    ):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)
    verified = Document(temporary)
    refreshed = [paragraph for paragraph in verified.paragraphs if paragraph.text.startswith(prefix)]
    if len(refreshed) != 1 or "构建无 WiX 或 ICE 警告" not in refreshed[0].text:
        raise AssertionError("packaging verification was not updated")
    if any(
        getattr(verified.core_properties, field)
        for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category")
    ):
        raise AssertionError("descriptive metadata remains")
    os.replace(temporary, report)
    temporary = None
finally:
    if temporary is not None:
        temporary.unlink(missing_ok=True)

print("Review report packaging verification updated.")
