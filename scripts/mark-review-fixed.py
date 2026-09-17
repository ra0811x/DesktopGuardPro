"""Mark verified review findings without changing their paragraph formatting."""
import argparse
import io
import os
from pathlib import Path
import tempfile
import zipfile

from docx import Document
from lxml import etree

parser = argparse.ArgumentParser()
parser.add_argument("findings", nargs="+")
args = parser.parse_args()
path = Path(__file__).resolve().parents[1] / "项目审查报告-2026-09-04.docx"
document = Document(path)
original = [(p.text, p._p.pPr.xml if p._p.pPr is not None else None) for p in document.paragraphs]
changed = set()
supplemental = {
    "S01": "升级检查存在竞争区间：",
    "S02": "自定义安装参数与服务固定路径不匹配：",
    "S03": "历史会话缺少可发现入口：",
    "S04": "未接入文件模块的补偿不完整：",
}
for finding in args.findings:
    prefix = supplemental.get(finding, finding + "  ")
    matches = [(i, p) for i, p in enumerate(document.paragraphs) if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one heading for {finding}")
    index, paragraph = matches[0]
    if "【已修复】" in paragraph.text:
        continue
    expected_runs = 2 if finding in supplemental else 1
    expected_style = "Normal" if finding in supplemental else "Heading 2"
    if paragraph.style.name != expected_style or len(paragraph.runs) != expected_runs:
        raise ValueError(f"Inspect formatting before changing {finding}")
    paragraph.runs[-1].text += "【已修复】"
    changed.add(index)
for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category", "identifier", "language", "content_status", "version"):
    setattr(document.core_properties, field, "")
buffer = io.BytesIO()
document.save(buffer)
temporary = None
try:
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".review-status-", suffix=".docx", delete=False) as output:
        temporary = Path(output.name)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root:
                    if etree.QName(node).localname in ("Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase"):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)
    verified = Document(temporary)
    assert len(verified.paragraphs) == len(original)
    for index, paragraph in enumerate(verified.paragraphs):
        text, properties = original[index]
        assert paragraph.text == text + ("【已修复】" if index in changed else "")
        assert (paragraph._p.pPr.xml if paragraph._p.pPr is not None else None) == properties
    assert not any(getattr(verified.core_properties, field) for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category"))
    with zipfile.ZipFile(temporary) as check:
        root = etree.fromstring(check.read("docProps/app.xml"))
        assert all(not node.text for node in root if etree.QName(node).localname in ("Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase"))
    os.replace(temporary, path)
    temporary = None
finally:
    if temporary is not None:
        temporary.unlink(missing_ok=True)
print("Marked verified findings: " + ", ".join(args.findings))
