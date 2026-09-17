"""Record verified Web repair batches in the existing Word review report."""

import io
import os
from pathlib import Path
import tempfile
import zipfile

from docx import Document
from lxml import etree


REPORT = Path(__file__).resolve().parents[1] / "项目审查报告-2026-09-04.docx"
REPAIR_SECTION = "Web 修复记录（2026 年 9 月 5 日）"
FORBIDDEN_PHRASES = ("建议", "总之", "现状", "下一步", "接下来", "不是")

STATUS_UPDATES = (
    (
        "W11  自动化测试未覆盖组件、浏览器与可访问性边界【待完善｜P3】",
        "W11  自动化测试未覆盖组件、浏览器与可访问性边界【部分修复｜P3】",
    ),
    (
        "W11  自动化测试未覆盖组件、浏览器与可访问性边界【部分修复｜P3】",
        "W11  自动化测试未覆盖组件、浏览器与可访问性边界【已修复｜P3】",
    ),
    (
        "W01  降级状态的主指标仍显示“正常”【待修复｜P1】",
        "W01  降级状态的主指标仍显示“正常”【已修复｜P1】",
    ),
    (
        "W06  风险摘要对数据来源的说明与实现不符【待修复｜P3】",
        "W06  风险摘要对数据来源的说明与实现不符【已修复｜P3】",
    ),
    (
        "W02  例行健康轮询短暂进入重连态并禁用操作【待修复｜P2】",
        "W02  例行健康轮询短暂进入重连态并禁用操作【已修复｜P2】",
    ),
    (
        "W03  开发模式 CSP 阻止 Vite 注入全部样式【待修复｜P2】",
        "W03  开发模式 CSP 阻止 Vite 注入全部样式【已修复｜P2】",
    ),
    (
        "W04  视图切换重复执行时间线查询与全会话风险计算【待修复｜P2】",
        "W04  视图切换重复执行时间线查询与全会话风险计算【已修复｜P2】",
    ),
    (
        "W05  时间线缺少有效后台动作的中文映射【待修复｜P2】",
        "W05  时间线缺少有效后台动作的中文映射【已修复｜P2】",
    ),
    (
        "W07  主导航缺少当前项语义，焦点外观对比不足【待修复｜P3】",
        "W07  主导航缺少当前项语义，焦点外观对比不足【已修复｜P3】",
    ),
    (
        "W08  历史会话选中行的辅助文字对比度偏低【待修复｜P3】",
        "W08  历史会话选中行的辅助文字对比度偏低【已修复｜P3】",
    ),
    (
        "W09  便携发布包缺少自定义图标，ICO 缺少小尺寸图层【待修复｜P2】",
        "W09  便携发布包缺少自定义图标，ICO 缺少小尺寸图层【部分修复｜P2】",
    ),
    (
        "W09  便携发布包缺少自定义图标，ICO 缺少小尺寸图层【部分修复｜P2】",
        "W09  便携发布包缺少自定义图标，ICO 缺少小尺寸图层【已修复｜P2】",
    ),
    (
        "W10  Web 发布产物缺少第三方许可告知【待修复｜P2】",
        "W10  Web 发布产物缺少第三方许可告知【部分修复｜P2】",
    ),
    (
        "W10  Web 发布产物缺少第三方许可告知【部分修复｜P2】",
        "W10  Web 发布产物缺少第三方许可告知【已修复｜P2】",
    ),
)

RECORDS = (
    (
        "批次 A：最小组件测试环境【已完成】",
        (
            "实现：frontend/package.json 与 package-lock.json 固定加入 happy-dom 20.14.0；frontend/src/SessionHistoryPanel.test.ts 使用 Vue createApp 挂载真实单文件组件，不引入额外组件测试框架。",
            "失败回归：依赖安装前运行组件测试，Vitest 因缺少 happy-dom 无法启动测试环境，稳定复现 W11 记录的组件测试基础缺口。",
            "通过结果：新增 2 项组件测试，覆盖区域可访问名称、查看按钮可访问名称、状态文字和 select 事件。全量结果为 8 个测试文件、32 项测试通过；vue-tsc、Vite 生产构建与 npm 离线依赖审计通过。",
            "保留范围：开发服务器 CSP 烟雾测试、主导航当前项和焦点检查、颜色对比自动检查仍由 W03、W07 与 W08 批次补齐。W11 保持“部分修复”。",
        ),
    ),
    (
        "批次 A2：真实 Edge 浏览器烟雾测试【已完成】",
        (
            "实现：frontend/browserSmoke.test.ts 由 Vitest 启动临时 Vite 开发服务器，并通过固定版本 playwright-core 1.63.0 驱动本机 Edge；测试结束后自动关闭浏览器与服务器。",
            "失败回归：项目依赖声明加入前，测试能够从外部环境加载浏览器驱动，但 package.json 中的固定版本断言收到空值，直接暴露测试不可独立复现的问题。",
            "通过结果：真实 Edge 自动检查开发 CSP、样式表加载、页面背景、应用网格布局、首页 aria-current 和控制台 CSP 错误。全量结果为 11 个测试文件、42 项测试通过；vue-tsc、Vite 生产构建与 npm 离线依赖审计通过，已知漏洞为 0。",
            "缺陷标记：W11 已形成组件挂载、真实浏览器、导航语义、焦点可见性和颜色对比度自动回归，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 B：采集健康与风险范围文案【已完成】",
        (
            "实现：frontend/src/serviceConnection.ts 显式返回后台原始健康状态；frontend/src/App.vue 根据连接与 degraded 状态显示“正常”“降级”或“恢复中”，风险摘要改为“基于完整会话事件”。",
            "失败回归：新增 frontend/src/App.test.ts 并挂载真实应用组件。修复前两项断言分别收到“正常”和“基于当前已加载会话事件”，与降级健康响应及完整会话风险实现冲突。",
            "通过结果：应用组件和连接状态定向测试共 14 项通过；全量结果为 9 个测试文件、34 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W01 与 W06 已满足各自验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 C：例行健康轮询保持在线【已完成】",
        (
            "实现：frontend/src/serviceConnection.ts 在健康检查开始时保留已有 online 阶段；首次连接与故障重试继续使用原连接状态，失败请求仍进入原有退避流程。",
            "失败回归：frontend/src/serviceConnection.test.ts 使用伪计时器触发例行轮询，并延迟第二次健康响应。修复前等待期间的阶段为 reconnecting，connected 同时变为 false。",
            "通过结果：连接状态定向测试 13 项通过；全量结果为 9 个测试文件、35 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W02 已满足例行成功轮询保持在线、失败请求继续重试的验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 D：开发期 CSP 样式加载【已完成】",
        (
            "实现：frontend/vite.config.ts 新增仅在 serve 模式启用的 CSP 转换插件，为 Vite 动态样式加入 'unsafe-inline'。frontend/index.html 保持生产策略 style-src 'self'。",
            "失败回归：新增 frontend/vite.config.test.ts，修复前无法找到开发期 CSP 插件，开发 HTML 也无法获得样式注入权限。",
            "通过结果：配置测试通过；全量结果为 10 个测试文件、36 项测试通过。vue-tsc 与 Vite 生产构建通过。真实 Edge 打开开发服务器后加载 7 张样式表，画布背景为 rgb(243, 247, 252)，应用外壳为网格布局，CSP 样式错误为 0。",
            "缺陷标记：W03 已满足开发样式可用和生产策略保持严格的验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 E：分析视图避免重复刷新【已完成】",
        (
            "实现：frontend/src/App.vue 移除普通分析标签切换中的 analysis.refresh 调用，保留会话 ID 变化监听、时间线显式刷新、筛选应用和加载更多入口。",
            "失败回归：frontend/src/App.test.ts 建立当前会话后依次切换时间线、风险分析和报告导出。修复前完整刷新共执行 4 次，其中 3 次来自同一会话的标签切换。",
            "通过结果：修复后会话变化触发 1 次完整加载，三个标签切换后调用次数仍为 1。应用与分析状态定向测试 7 项通过；全量结果为 10 个测试文件、37 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W04 已满足同一会话普通标签切换不重复完整刷新的验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 F：后台动作中文映射【已完成】",
        (
            "实现：frontend/src/timelinePresentation.ts 为 service_recovered_after_interruption、file_rename_old_name 和 file_rename_new_name 增加中文标签。未知扩展动作继续使用可读的空格英文回退。",
            "失败回归：frontend/src/timelinePresentation.test.ts 先断言中断恢复事件和未配对重命名事件的中文输出。修复前中断恢复事件实际返回 service recovered after interruption。",
            "通过结果：三种后台动作的中文断言通过；表现层定向测试 3 项通过，全量结果为 10 个测试文件、37 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W05 已覆盖审查记录中的后台动作，并补充同类新名称分支，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 G：主导航语义与焦点外观【已完成】",
        (
            "实现：frontend/src/App.vue 为五个导航按钮按当前视图设置 aria-current=page；frontend/src/styles.css 为深蓝侧栏中的导航焦点设置不透明白色三像素外环。",
            "失败回归：frontend/src/App.test.ts 挂载应用后读取首页和历史会话按钮，修复前 aria-current 为空；同一测试读取导航焦点 CSS，修复前没有能够达到 3:1 的侧栏专用焦点环。",
            "通过结果：当前页面语义会随视图切换；测试自动计算白色焦点环与 #18324d 侧栏背景的对比度并通过 3:1 门槛。应用定向测试 5 项通过，全量结果为 10 个测试文件、39 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W07 已满足读屏当前项和键盘焦点外观验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 H：历史选中行文字对比度【已完成】",
        (
            "实现：frontend/src/SessionHistoryPanel.vue 为选中行的 12 像素会话 ID 使用已有 --text-secondary 文字色，普通行继续使用 --text-muted。",
            "失败回归：frontend/src/SessionHistoryPanel.test.ts 挂载选中会话，确认 selected 状态，再读取组件规则和全局颜色令牌。修复前没有选中行专用文字色，原组合测得约 4.09:1。",
            "通过结果：测试自动计算 #314f6b 与 #dce7f0 的对比度约为 6.80:1，高于 4.5:1 门槛。组件定向测试 3 项通过，全量结果为 10 个测试文件、40 项测试通过。vue-tsc 与 Vite 生产构建通过。",
            "缺陷标记：W08 已满足选中状态和小字号文字对比度验收条件，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 I1：便携包图标交付【已完成】",
        (
            "实现：cmd/desktop-guard-release/main.go 在构建目录中复制 desktop-guard-pro.ico，把发布清单升级为 schema v2，并在 assets 字段记录图标大小与 SHA-256。ZIP 文件列表加入图标；签名循环仍只验证四个可执行组件。",
            "失败回归：cmd/desktop-guard-release/main_test.go 在实现前无法找到图标文件常量、清单资产字段和完整发布文件列表，编译失败并直接暴露发布链缺口。",
            "通过结果：发布命令定向测试通过，ZIP 条目和清单资产断言通过；go test ./... 全项目通过。",
            "保留范围：imagegen 两次输出均没有 Alpha 通道，并保留棋盘格或渐变，未写入项目。多尺寸、无文字、透明背景 ICO 仍待生成和小尺寸验收；W09 保持“部分修复”。",
        ),
    ),
    (
        "批次 I2：多尺寸透明应用图标与 1.0.5 发布包【已完成】",
        (
            "实现：assets/desktop-guard-pro.ico 保留盾牌与显示器主体，移除小尺寸下不可辨识的文字，封装 16、24、32、48、64、128、256 七个 32 位 PNG 图层。每层四角保持透明。1.0.5 便携目录、ZIP 和 MSI 均使用该图标。",
            "失败回归：cmd/desktop-guard-release/main_test.go 首次运行报告图标只有一个图层；首次多尺寸封装又报告 16×16 图层左上角透明度异常。按尺寸独立缩放并校正透明角后，同一用例通过。",
            "通过结果：go test ./... 全项目通过；前端 13 个测试文件、49 项测试通过，类型检查与生产构建通过。ZIP 含 7 个预期条目，4 个可执行组件签名有效；MSI 签名、WiX/ICE、事务顺序和发布门禁检查通过。",
            "缺陷标记：W09 已覆盖便携包图标交付、多尺寸图层、透明背景和小尺寸自动验收，标题更新为“已修复”。",
        ),
    ),
    (
        "批次 J1：第三方许可告知进入便携包【已完成】",
        (
            "实现：根目录新增 THIRD-PARTY-NOTICES.txt，记录 Vue 3.5.41 与 Wails Runtime 3.0.0-beta.12 的项目地址、版权和完整 MIT 条款。发布工具把许可文件作为独立资产写入 schema v2 清单与 ZIP。",
            "失败回归：发布测试先要求第二个清单资产，修复前缺少许可文件常量和对应资产，测试无法编译。",
            "通过结果：发布命令定向测试、go test ./... 与 go vet ./... 通过；清单资产顺序、文件名和 ZIP 条目均被自动检查。",
            "保留范围：MSI 尚未安装许可文件，批次 J2 继续处理安装器来源和包表检查；W10 保持“部分修复”。",
        ),
    ),
    (
        "批次 J2：第三方许可告知进入 MSI【已完成】",
        (
            "实现：installer/DesktopGuardPro.wxs 把 THIRD-PARTY-NOTICES.txt 加入 ApplicationFiles 组件，安装与卸载均由产品文件生命周期管理。",
            "失败回归：scripts/test-msi-package.ps1 先要求 MSI File 表包含 ThirdPartyNotices。修复前临时 MSI 可编译，但文件表检查明确报告许可文件缺失。",
            "通过结果：临时 MSI 编译无 WiX 或 ICE 警告，文件表包含许可告知；提升动作、回滚顺序、四个可执行文件哈希、签名门禁、回执和失败发布保护检查通过。本轮没有执行产品安装。",
            "缺陷标记：W10 已同时覆盖便携 ZIP、schema v2 发布清单和 MSI，标题更新为“已修复”。",
        ),
    ),
    (
        "W12  MSI 安装缺少标准交互并显示维护终端窗口【已修复｜P2】",
        (
            "实现：cmd/desktop-guard-release/main.go 将安装器调用的维护程序编译为 Windows GUI 子系统。installer/DesktopGuardPro.wxs 加入中文安装目录页、桌面快捷方式选项、确认与进度流程，并在完成页提供启动 Desktop Guard Pro 选项。桌面快捷方式由独立 MSI Feature 管理，可参与回滚、升级和卸载。",
            "失败回归：发布测试先报告维护程序未启用 GUI 子系统；安装器结构测试随后报告安装目录与桌面选项缺失；WiX 包测试在扩展接入前报告 WixCA 和标准对话框引用无法解析。已签名发布目录复用测试同时复现图标资产未复制的问题。",
            "通过结果：go test ./... 全项目通过，前端 13 个测试文件、49 项测试通过。WiX/ICE 编译无警告，MSI 数据表确认选项对话框、桌面 Feature、AddLocal、Remove 和完成页启动事件存在。1.0.6 中的维护程序与主界面均为 Windows GUI 子系统，MSI 和四个组件签名有效。",
            "验证边界：自动检查覆盖安装器结构、包表、静默安装不触发完成页动作、签名和发布清单。本轮未执行真实安装与卸载，避免改动本机服务和受保护会话。",
        ),
    ),
    (
        "批次 K2：安装器扁平品牌视觉【已完成】",
        (
            "实现：assets/installer-dialog.bmp 使用 493×312 白色画布、164 像素深蓝侧栏和小尺寸产品图标；assets/installer-banner.bmp 使用 493×58 浅蓝表面、浅蓝边框和右侧产品图标。installer/DesktopGuardPro.wxs 通过 WixUIDialogBmp 与 WixUIBannerBmp 嵌入两项资源。",
            "失败回归：提取首次生成的 1.0.6 MSI 后，WixUI_Bmp_Dialog 与 WixUI_Bmp_Banner 仍为 WiX 默认红黑高饱和图形。安装器资源测试在自定义位图与变量缺失时明确失败。",
            "通过结果：资源测试确认两张位图尺寸、深蓝、白色、浅蓝和边框关键像素；WiX/ICE 编译无警告。重新生成的 MSI 中，两项 Binary 流的 SHA-256 与项目资源一致，签名状态有效。",
            "视觉约束：安装器使用扁平、低饱和的蓝、白、灰配色，不含毛玻璃、渐变色块和大型装饰图形。",
        ),
    ),
    (
        "W13  GUI 维护程序将输出句柄错误误判为 MSI 操作失败【已修复｜P1】",
        (
            "实现：cmd/desktop-guard-maintenance/main.go 保留维护操作的真实执行结果。Windows GUI 子系统没有可用标准输出句柄时，JSON 回执写入失败不再把已经完成的 prepare、apply 或 commit 阶段改写为失败；维护操作自身返回的错误仍按原有退出码处理。",
            "失败回归：cmd/desktop-guard-maintenance/main_test.go 新增不可用输出句柄场景。修复前，模拟输出写入失败后退出码为 1，能够复现 Windows Installer 中 PrepareDesktopGuard 动作触发 1722 并回滚的问题。",
            "通过结果：修复后同一用例退出码为 0。go test ./...、go vet ./...、前端 13 个测试文件与 49 项测试、类型检查、生产构建、MSI 命令及包表测试全部通过。1.0.7 GUI 维护程序执行 diagnose 的退出码为 0。",
            "真实升级：在本机将已运行的 1.0.5 静默升级到 1.0.7，Windows Installer 返回 0，产品版本更新为 1.0.7。DesktopGuardPro 服务保持自动启动和运行状态，安装清单记录 1.0.7，四个核心组件的文件哈希与签名均通过，维护事务日志已清除，安装过程没有启动前端程序。",
        ),
    ),
    (
        "W14  同一 MSI 重装与已保留恢复目录发生冲突【已修复｜P1】",
        (
            "实现：internal/maintenance/msi_data_windows.go 为每次 MSI 事务创建随机唯一的恢复目录。internal/maintenance/msi_transaction_windows.go 把目录名写入受保护事务日志，并在读取时校验目录名与事务标识，回滚和提交始终使用本次事务对应的目录。旧格式事务没有目录字段时仍按原路径读取。",
            "失败回归：internal/maintenance/msi_data_windows_test.go 使用相同 ProductCode 连续执行两次数据备份。修复前第二次备份在 os.Mkdir 创建同名目录时返回目录已存在，稳定复现 PrepareDesktopGuard 错误 1722。",
            "通过结果：回归用例确认两次备份目录不同，第二次备份内容正确。go test ./...、go vet ./...、前端 13 个测试文件与 49 项测试、类型检查、生产构建及 MSI 测试全部通过。",
            "真实重装：本机使用 1.0.8 完成安装、卸载并保留数据、再次安装同一 MSI。三次 Windows Installer 操作均返回 0，两次安装生成不同恢复目录。产品版本、安装清单均为 1.0.8，服务保持自动启动和运行状态，四个核心组件的哈希与签名通过，事务日志已清除，安装日志没有 Return value 3。",
        ),
    ),
    (
        "W15  MSI 卸载遗留重启删除登记会影响同路径重装【已修复｜P1】",
        (
            "实现：MSI 卸载动作通过 --msi-managed 标记交由 Windows Installer 删除程序文件，不再调用独立卸载流程的延迟删除登记，卸载结果也不再报告需要重启。安装阶段在签名和发布清单验证通过后，清除精确指向 Desktop Guard Pro 四个核心程序及安装目录的旧删除记录。",
            "失败回归：卸载测试先要求 MSI 托管模式跳过 schedule_program_removal；维护命令测试先要求识别 --msi-managed；安装器测试先要求 RemoveDesktopGuard 传递该标记。旧记录过滤测试同时覆盖目标删除项、其他软件删除项和文件重命名项。修复前前三项缺少对应行为，过滤函数也不存在。",
            "通过结果：定向测试确认 MSI 托管卸载不登记重启删除，独立卸载行为保持原样；旧记录过滤只移除目标为空且路径精确匹配的删除对。go test ./...、go vet ./...、前端 13 个测试文件与 49 项测试、类型检查、生产构建及 MSI 测试全部通过。",
            "真实验证：本机从 1.0.8 升级到 1.0.9 时清除 10 条历史待删除记录。随后完成 1.0.9 卸载并保留数据、再次安装同一 MSI，Windows Installer 均返回 0，卸载后和重装后待删除记录均为 0。当前版本与安装清单均为 1.0.9，服务保持自动启动和运行状态，诊断退出码为 0，四个核心组件的哈希与签名通过，事务日志已清除。",
        ),
    ),
)


def find_exact(document, text):
    return [paragraph for paragraph in document.paragraphs if paragraph.text == text]


document = Document(REPORT)
original_paragraph_count = len(document.paragraphs)
original_table_count = len(document.tables)
original_sections = [section._sectPr.xml for section in document.sections]
original_paragraphs = [
    (paragraph.text, paragraph.style.name)
    for paragraph in document.paragraphs
]

replacements = {}
terminal_statuses = {}
for pending, completed in STATUS_UPDATES:
    terminal_statuses[pending.split("  ", 1)[0]] = completed
for pending, completed in STATUS_UPDATES:
    finding_id = pending.split("  ", 1)[0]
    if find_exact(document, terminal_statuses[finding_id]):
        continue
    pending_matches = find_exact(document, pending)
    completed_matches = find_exact(document, completed)
    if len(completed_matches) == 1 and not pending_matches:
        continue
    if len(pending_matches) != 1 or completed_matches:
        raise ValueError(f"Inspect finding status before updating {pending!r}")
    paragraph = pending_matches[0]
    if len(paragraph.runs) != 1:
        raise ValueError(f"Inspect finding heading formatting before replacing {pending!r}")
    paragraph.runs[0].text = completed
    replacements[pending] = completed

missing_records = [record for record in RECORDS if not find_exact(document, record[0])]
if not replacements and not missing_records:
    print("All configured Web repair records already exist.")
    raise SystemExit(0)

if missing_records and not find_exact(document, REPAIR_SECTION):
    document.add_heading(REPAIR_SECTION, level=2)
    document.add_paragraph(
        "本节记录 Web 专项发现的代码处理和回归结果。只有完整满足验收条件的条目才会标为“已修复”；仍有覆盖范围待补齐的条目使用“部分修复”。"
    )

for heading, paragraphs in missing_records:
    document.add_heading(heading, level=3)
    for text in paragraphs:
        document.add_paragraph(text)

for field in (
    "title", "subject", "author", "last_modified_by", "comments", "keywords",
    "category", "identifier", "language", "content_status", "version",
):
    setattr(document.core_properties, field, "")

buffer = io.BytesIO()
document.save(buffer)
temporary = None
try:
    with tempfile.NamedTemporaryFile(
        dir=REPORT.parent, prefix=".web-fix-", suffix=".docx", delete=False,
    ) as output:
        temporary = Path(output.name)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root:
                    if etree.QName(node).localname in (
                        "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
                    ):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)

    verified = Document(temporary)
    if len(verified.paragraphs) < original_paragraph_count:
        raise AssertionError("existing report content was removed")
    if len(verified.tables) != original_table_count:
        raise AssertionError("existing tables changed")
    if [section._sectPr.xml for section in verified.sections] != original_sections:
        raise AssertionError("page layout changed")
    for index, (old_text, old_style) in enumerate(original_paragraphs):
        paragraph = verified.paragraphs[index]
        expected_text = replacements.get(old_text, old_text)
        if paragraph.text != expected_text or paragraph.style.name != old_style:
            raise AssertionError(f"unexpected change at paragraph {index}")
    appended = "\n".join(paragraph.text for paragraph in verified.paragraphs[original_paragraph_count:])
    if any(phrase in appended for phrase in FORBIDDEN_PHRASES):
        raise AssertionError("repair record contains a prohibited phrase")
    for completed in set(terminal_statuses.values()):
        if len(find_exact(verified, completed)) != 1:
            raise AssertionError(f"finding status is missing: {completed}")
    for heading, _ in RECORDS:
        if len(find_exact(verified, heading)) != 1:
            raise AssertionError(f"repair record is missing: {heading}")
    if any(
        getattr(verified.core_properties, field)
        for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category")
    ):
        raise AssertionError("descriptive metadata remains")
    with zipfile.ZipFile(temporary) as archive:
        root = etree.fromstring(archive.read("docProps/app.xml"))
        if archive.testzip() is not None:
            raise AssertionError("report package CRC check failed")
        if any(
            node.text for node in root
            if etree.QName(node).localname in (
                "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
            )
        ):
            raise AssertionError("application metadata remains")
    os.replace(temporary, REPORT)
    temporary = None
finally:
    if temporary is not None:
        temporary.unlink(missing_ok=True)

print("Configured Web repair records appended and verified.")
