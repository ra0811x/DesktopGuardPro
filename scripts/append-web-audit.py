"""Append the verified Web application audit while preserving the existing report."""

import io
import os
from pathlib import Path
import tempfile
import zipfile

from docx import Document
from lxml import etree


REPORT = Path(__file__).resolve().parents[1] / "项目审查报告-2026-09-04.docx"
SECTION_TITLE = "Web 应用与图形资源专项复核（2026 年 9 月 5 日）"
FORBIDDEN_PHRASES = ("建议", "总之", "现状", "下一步", "接下来", "不是")
BATCHES = (
    "批次 A：W11，先建立最小组件测试环境；依赖清单、锁文件与测试配置合计不超过三个文件。",
    "批次 B：W01 与 W06，校正首页状态和风险范围文案，在已有应用组件测试中添加回归。",
    "批次 C：W02，保持健康例行轮询期间的在线状态，扩展 serviceConnection 单元测试。",
    "批次 D：W03，区分开发期和生产期 CSP，增加开发服务器样式烟雾测试。",
    "批次 E：W04，拆分时间线与风险数据刷新，增加视图切换调用次数回归。",
    "批次 F：W05，补全已知事件动作本地化及表现层测试。",
    "批次 G：W07，修复导航语义与焦点外观并补测试；批次 H：W08，单独修复历史选中行对比度并补测试。",
    "批次 I：W09，处理便携图标、ICO 图层与发布测试；批次 J1：W10 先生成许可文件并纳入 ZIP；批次 J2：把许可文件纳入 MSI 并补安装器检查。",
)


FINDINGS = (
    (
        "W01  降级状态的主指标仍显示“正常”【待修复｜P1】",
        "定位：frontend/src/App.vue 第 271 行只按服务连接布尔值输出“正常”或“恢复中”，未读取已经计算出的 degraded 健康状态。",
        "复现与影响：模拟 GetHealth 返回 degraded 后，页眉显示“采集服务降级”，首页“核心保护”主值同时显示“正常”。同一页面给出相互冲突的保护结论，可能使用户忽略采集缺口。",
        "验收用例：挂载应用并让桥接层返回 degraded，断言核心保护主值、页眉和状态说明一致表达降级；online、offline、reconnecting 分别覆盖。",
    ),
    (
        "W02  例行健康轮询短暂进入重连态并禁用操作【待修复｜P2】",
        "定位：frontend/src/serviceConnection.ts 第 113—126 行在每次非首次健康检查开始时无条件写入 reconnecting；第 69 行据此把 connected 置为 false，第 225—231 行每 10 秒触发一次例行检查。",
        "复现与影响：缩短轮询周期并延迟第二次成功响应，界面会显示“正在重连”，主操作在请求完成前处于禁用状态。健康服务也会周期性闪烁故障提示，慢设备上的干扰更明显。",
        "验收用例：使用伪计时器和可控 Promise 触发一次成功例行轮询，响应等待期间仍保持 online、connected 与操作可用；失败请求仍进入重试流程。",
    ),
    (
        "W03  开发模式 CSP 阻止 Vite 注入全部样式【待修复｜P2】",
        "定位：frontend/index.html 第 6—7 行固定声明 style-src 'self'。Vite 开发服务器通过内联 style 元素注入组件与全局样式，该声明未提供开发期所需的放行方式。",
        "复现与影响：启动 npm run dev 并用 Edge 打开页面，浏览器报告 CSP 拒绝内联样式，document.styleSheets 为空，页面退化为无样式内容。生产构建提取为同源 CSS 文件，未复现该问题。",
        "验收用例：增加开发服务器浏览器烟雾测试，断言至少加载一张样式表、应用外壳获得预期布局；生产产物继续断言 CSP 不放宽脚本执行策略。",
    ),
    (
        "W04  视图切换重复执行时间线查询与全会话风险计算【待修复｜P2】",
        "定位：frontend/src/App.vue 第 137—143 行在进入时间线、风险或报告视图时调用 analysis.refresh；第 170—173 行又监听会话变化。internal/service/analysis_api.go 第 258 行的风险计算读取完整会话事件。",
        "复现与影响：在模拟桥接层中依次切换时间线、风险和报告，每次均观察到 QueryTimeline 与 EvaluateRisk 成对调用。大体量会话会反复扫描完整事件集合，增加数据库读取、界面等待和资源占用。",
        "验收用例：记录桥接调用次数，同一会话的普通标签切换不得重复完整刷新；显式刷新、筛选变化和会话变化仍按各自数据依赖发起调用。",
    ),
    (
        "W05  时间线缺少有效后台动作的中文映射【待修复｜P2】",
        "定位：cmd/desktop-guard-service/runtime.go 第 168 行会产生 service_recovered_after_interruption；internal/collector/file.go 第 160 行可产生 file_rename_old_name。frontend/src/timelinePresentation.ts 第 24—38 行的动作表未收录二者，第 45 行回退为带空格的英文。",
        "复现与影响：恢复事件在中文界面显示为“service recovered after interruption”；跨通知批次中未配对的旧文件名事件也会落入英文回退。关键采集边界缺少一致的中文表达。",
        "验收用例：为两种后台动作增加表现层单元测试，并遍历后台已知动作集合，确保中文界面不会对已知动作使用英文回退。",
    ),
    (
        "W06  风险摘要对数据来源的说明与实现不符【待修复｜P3】",
        "定位：frontend/src/App.vue 第 277 行写明“基于当前已加载会话事件”；internal/service/analysis_api.go 第 257—260 行实际读取完整会话后计算风险。",
        "复现与影响：用户可能把风险结果理解为仅覆盖当前时间线页面。此前已经交付的完整会话风险语义没有在界面中准确呈现。",
        "验收用例：组件测试断言存在风险结果时显示“完整会话”范围，且翻页和时间线筛选不改变该说明。",
    ),
    (
        "W07  主导航缺少当前项语义，焦点外观对比不足【待修复｜P3】",
        "定位：frontend/src/App.vue 第 198—202 行只用 active 类表达当前视图，未提供 aria-current 或等价状态。frontend/src/styles.css 第 56—63 行在深蓝侧栏上使用低透明度蓝色焦点环。",
        "复现与影响：键盘导航时，焦点环与侧栏背景的测得对比度约为 1.15:1，边框约为 2.18:1，低于 3:1 的焦点外观基准；读屏也无法直接识别当前页面。",
        "验收用例：逐项 Tab 导航并断言当前按钮具有 aria-current，焦点指示器在侧栏背景上的对比度达到 3:1，鼠标与键盘状态均可辨认。",
    ),
    (
        "W08  历史会话选中行的辅助文字对比度偏低【待修复｜P3】",
        "定位：frontend/src/SessionHistoryPanel.vue 第 48—50 行把 12 像素会话 ID 置于选中背景 #dce7f0 上，文字沿用 #587089。",
        "复现与影响：测得对比度约为 4.09:1，低于普通小字号文本 4.5:1 的基准。会话 ID 在选中行中比未选中行更难辨认。",
        "验收用例：覆盖普通、悬停、选中与键盘焦点状态，自动计算会话 ID 的前景和背景对比度，并在 100% 与 200% 缩放下人工核对。",
    ),
    (
        "W09  便携发布包缺少自定义图标，ICO 缺少小尺寸图层【待修复｜P2】",
        "定位：cmd/desktop-guard-ui/main.go 第 108—117 行只从可执行文件同目录读取 desktop-guard-pro.ico，读取失败即使用框架默认图标。cmd/desktop-guard-release/main.go 第 226—231 行只返回四个可执行文件名，当前 ZIP 也仅含四个 EXE 与 release-manifest.json。MSI 已单独安装该图标。",
        "复现与影响：解包 DesktopGuardPro-1.0.4-windows-amd64.zip 后没有 ICO，便携运行会回退到默认图标。现有 ICO 结构有效，但只含一个 256×256 图层；缩放至 16、24、32 像素时产品文字不可读，系统需要从单一大图降采样。",
        "验收用例：发布构建测试断言 ZIP 与清单包含 ICO；在无安装环境启动便携版，核对窗口、任务栏和托盘图标；结构测试要求包含 16、24、32、48、256 像素图层。",
    ),
    (
        "W10  Web 发布产物缺少第三方许可告知【待修复｜P2】",
        "定位：生产包会分发 Vue 与 Wails Web Runtime 的打包代码，两者采用 MIT 许可。仓库根目录、frontend/dist 与当前发布 ZIP 中未找到 LICENSE、NOTICE 或 THIRD-PARTY-NOTICES 文件，压缩后的脚本也未保留许可注释。",
        "复现与影响：面向用户分发时缺少依赖许可文本和版权归属，形成发布合规缺口。",
        "验收用例：生成第三方许可清单并纳入发布清单、ZIP 与 MSI；构建测试校验文件存在且至少覆盖全部生产依赖，版本变化时同步更新。",
    ),
    (
        "W11  自动化测试未覆盖组件、浏览器与可访问性边界【待完善｜P3】",
        "定位：当前 7 个测试文件共 30 项测试，均未导入 .vue 组件，也未安装组件挂载或浏览器测试环境；项目没有开发期 CSP、键盘导航、颜色对比和视图调用次数的自动检查。",
        "复现与影响：本轮确认的状态矛盾、轮询闪烁、重复分析、导航语义与开发样式故障均未被已有测试捕获。",
        "验收用例：按 W01—W08 的复现条件补齐组件与浏览器测试，并保留现有 30 项单元测试作为基础回归。",
    ),
)


def add_text_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    return paragraph


document = Document(REPORT)
if any(paragraph.text == SECTION_TITLE for paragraph in document.paragraphs):
    batch_paragraphs = [
        paragraph for paragraph in document.paragraphs
        if paragraph.text.startswith("批次 ")
    ]
    if tuple(paragraph.text for paragraph in batch_paragraphs) == BATCHES:
        print("Web audit section already exists.")
        raise SystemExit(0)
    if len(batch_paragraphs) != len(BATCHES):
        raise ValueError("Inspect the existing Web audit batch plan before updating it.")
    original_styles = [paragraph.style.name for paragraph in batch_paragraphs]
    for paragraph, value in zip(batch_paragraphs, BATCHES):
        paragraph.text = value
    temporary = REPORT.with_name(".web-audit-plan.docx")
    try:
        document.save(temporary)
        verified = Document(temporary)
        verified_batches = [
            paragraph for paragraph in verified.paragraphs
            if paragraph.text.startswith("批次 ")
        ]
        if tuple(paragraph.text for paragraph in verified_batches) != BATCHES:
            raise AssertionError("batch plan update failed")
        if [paragraph.style.name for paragraph in verified_batches] != original_styles:
            raise AssertionError("batch plan paragraph styles changed")
        os.replace(temporary, REPORT)
    finally:
        temporary.unlink(missing_ok=True)
    print("Web audit batch plan updated and verified.")
    raise SystemExit(0)

original_paragraph_count = len(document.paragraphs)
original_table_count = len(document.tables)
original_section_xml = [section._sectPr.xml for section in document.sections]
original_body_xml = [
    etree.tostring(child)
    for child in document.element.body
    if etree.QName(child).localname != "sectPr"
]

document.add_heading(SECTION_TITLE, level=1)
add_text_paragraph(
    document,
    "本节复核 Wails 桌面应用中的 Web 前端、图形资源与发布链路。审查覆盖源码、生产构建、开发服务器、模拟桌面桥接的浏览器交互、响应式布局、键盘与颜色可访问性、依赖公告、发布 ZIP 和 ICO 结构。仓库未包含 Dockerfile、Compose、OCI 或其他容器镜像定义，容器镜像没有可执行的审查对象。",
)

document.add_heading("审查结论", level=2)
add_text_paragraph(
    document,
    "本轮确认 11 项待处理内容：1 项 P1、6 项 P2、3 项 P3 和 1 项 P3 质量缺口。P1 涉及降级状态被主指标显示为正常；P2 集中在轮询状态、开发样式、重复分析、事件本地化、便携图标和许可告知。生产界面的视觉基线、构建、已有单元测试、离线依赖检查和页面级响应式布局均通过。",
)

document.add_heading("发现清单", level=2)
add_text_paragraph(
    document,
    "下列条目保存复现路径、用户影响与修复验收条件。全部条目保持“待修复”或“待完善”，未修改任何 Web 源码与发布产物。",
)
for title, location, impact, acceptance in FINDINGS:
    document.add_heading(title, level=3)
    add_text_paragraph(document, location)
    add_text_paragraph(document, impact)
    add_text_paragraph(document, acceptance)

document.add_heading("通过项", level=2)
add_text_paragraph(
    document,
    "以下结果用于限定问题范围，也作为后续回归基线。通过结果不覆盖真实 Windows 安装、原生 Wails 窗口与服务故障注入。",
)
for item in (
    "构建与测试：npm run build 通过，TypeScript 检查与 Vite 生产构建完成；7 个测试文件中的 30 项测试全部通过。生产产物没有 source map，也没有外部网络资源。",
    "响应式与视觉：在 320、560、768、992、1280 像素宽度下未发现页面级横向溢出，宽表格在自身容器内滚动。界面采用深蓝、白、灰蓝的克制配色，未使用渐变、毛玻璃或大面积装饰图。",
    "浏览器交互：使用模拟 Wails 桥接覆盖首页、历史会话、时间线、风险分析和报告导出。表单控件具有可访问名称，标题层级可辨认，生产模式控制台仅出现脱离原生环境运行时的预期提示。",
    "前端安全：源码未发现 v-html、innerHTML、eval、new Function、document.write、动态远程资源或浏览器存储敏感数据。报告文件名经过 Windows 字符清理，HTML 与 Markdown 转义已有后端回归。",
    "依赖检查：npm audit --offline 报告 0 个已知漏洞。Vite 8.2.2 已高于 GHSA-fx2h-pf6j-xcff 的 8.0.16 修复版本；Vitest 4.1.11 已覆盖 GHSA-82fw-gwwq-j7x9 与 GHSA-p63j-vcc4-9vmv 的修复版本。",
    "Vue 监视项：项目依赖树包含 @vue/server-renderer 3.5.41，对应 GHSA-g2v6-rqmx-r4w6 的受影响范围；当前应用只使用客户端 createApp，源码与生产包均未调用 SSR 渲染路径，因此该公告在当前运行路径中不可达。",
):
    add_text_paragraph(document, item, style="List Bullet")

document.add_heading("图形资源与容器镜像边界", level=2)
add_text_paragraph(
    document,
    "仓库内唯一独立图形资源为 assets/desktop-guard-pro.ico。文件签名为有效 ICO，包含一个 256×256 图层，SHA-256 为 803494FB467A6F325FC1A1717CA2A0188A533B6A79D24B08F09D372A20B22FA7。MSI 源文件会安装并引用它，便携 ZIP 未携带它，相关问题记录为 W09。",
)
add_text_paragraph(
    document,
    "仓库没有容器构建文件、基础镜像声明、镜像标签、软件物料清单或镜像签名流程。Desktop Guard Pro 当前以 Windows EXE、MSI 与便携 ZIP 交付；容器漏洞扫描、Dockerfile 最小化与 OCI 签名在本项目当前交付形态下不适用。",
)

document.add_heading("审查限制与交付边界", level=2)
add_text_paragraph(
    document,
    "下列限制单独记录，防止把已覆盖范围扩展到未执行环境。限制不会改变已复现问题的判定。",
)
for item in (
    "在线 npm 漏洞库查询会向 npm 发送依赖树摘要，本轮没有获得该外部传输授权，因此只执行本地缓存审计，并用官方安全公告交叉核对直接工具链。",
    "浏览器功能复核使用模拟 Wails 桥接；未在本轮启动真实 Windows 服务、原生 WebView、Windows Hello 窗口或执行 MSI 安装。",
    "当前 1.0.4 签名 EXE、MSI 与 ZIP 早于 2026 年 9 月 5 日的源码和前端构建，未包含最新源码修复。该边界已经在原报告的修复记录中说明，本节不重复登记为新缺陷。",
    "Wails Go 模块与前端 Runtime 均为 3.0.0-beta.12，版本保持一致。官方仍将 v3 标为预发布；稳定性风险需在真实 WebView 和安装环境中持续验证。",
):
    add_text_paragraph(document, item, style="List Bullet")

document.add_heading("修复批次划分", level=2)
add_text_paragraph(
    document,
    "代码修复可拆为小批次，每批严格控制在三个文件以内，并先加入失败回归再修改实现。执行前仍需取得用户对该批计划的确认。",
)
for item in BATCHES:
    add_text_paragraph(document, item, style="List Bullet")

for field in (
    "title", "subject", "author", "last_modified_by", "comments", "keywords",
    "category", "identifier", "language", "content_status", "version",
):
    setattr(document.core_properties, field, "")

buffer = io.BytesIO()
document.save(buffer)
temporary = None
try:
    with tempfile.NamedTemporaryFile(
        dir=REPORT.parent, prefix=".web-audit-", suffix=".docx", delete=False,
    ) as output:
        temporary = Path(output.name)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root:
                    if etree.QName(node).localname in (
                        "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
                    ):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)

    verified = Document(temporary)
    if len(verified.paragraphs) <= original_paragraph_count:
        raise AssertionError("report was not extended")
    if len(verified.tables) != original_table_count:
        raise AssertionError("existing tables changed")
    if [section._sectPr.xml for section in verified.sections] != original_section_xml:
        raise AssertionError("page layout changed")
    verified_body_xml = [
        etree.tostring(child)
        for child in verified.element.body
        if etree.QName(child).localname != "sectPr"
    ]
    if verified_body_xml[:len(original_body_xml)] != original_body_xml:
        raise AssertionError("existing document content or formatting changed")
    appended = "\n".join(paragraph.text for paragraph in verified.paragraphs[original_paragraph_count:])
    if any(phrase in appended for phrase in FORBIDDEN_PHRASES):
        raise AssertionError("appended text contains a prohibited phrase")
    if sum("【待修复" in paragraph.text for paragraph in verified.paragraphs) < 10:
        raise AssertionError("one or more Web findings are missing")
    if not any(paragraph.text == SECTION_TITLE for paragraph in verified.paragraphs):
        raise AssertionError("Web audit heading is missing")
    if any(
        getattr(verified.core_properties, field)
        for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category")
    ):
        raise AssertionError("descriptive metadata remains")
    with zipfile.ZipFile(temporary) as archive:
        root = etree.fromstring(archive.read("docProps/app.xml"))
        if any(
            node.text for node in root
            if etree.QName(node).localname in (
                "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
            )
        ):
            raise AssertionError("application metadata remains")
    os.replace(temporary, REPORT)
    temporary = None
finally:
    if temporary is not None:
        temporary.unlink(missing_ok=True)

print("Web application audit appended and verified.")
