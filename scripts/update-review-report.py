"""Append verified repair records while preserving the review evidence."""

import io
import os
from pathlib import Path
import re
import tempfile
import zipfile

from docx import Document
from lxml import etree


REPORT = Path(__file__).resolve().parents[1] / "项目审查报告-2026-09-04.docx"
FIXED_IDS = (
    "R01", "R02", "R03", "R04", "R06", "R08", "R09", "R10", "R11",
    "R12", "R13", "R14", "R15", "R16", "R17", "R18", "R19", "R20",
    "R21", "R22", "R23", "R24", "R25", "R26", "R27",
)
SUPPLEMENTAL_PREFIXES = (
    "升级检查存在竞争区间：",
    "自定义安装参数与服务固定路径不匹配：",
    "历史会话缺少可发现入口：",
    "未接入文件模块的补偿不完整：",
)


def set_single_run(paragraph, value):
    if len(paragraph.runs) != 1:
        raise ValueError(f"Inspect formatting before replacing: {paragraph.text}")
    paragraph.runs[0].text = value


def find_one(document, prefix):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}")
    return matches[0]


document = Document(REPORT)
if any(paragraph.text.startswith("修复与复核记录（截至") for paragraph in document.paragraphs):
    print("Review report repair records already exist.")
    raise SystemExit(0)
original_count = len(document.paragraphs)
original_tables = len(document.tables)
original_sections = [section._sectPr.xml for section in document.sections]
original_properties = {
    paragraph.text: paragraph._p.pPr.xml if paragraph._p.pPr is not None else None
    for paragraph in document.paragraphs
}

set_single_run(
    document.paragraphs[2],
    "本报告保存 2026 年 9 月 4 日审查时的证据、触发方式和风险判断。"
    "随后按小任务完成代码修复与回归；正文中的定位仍指向修复前快照，文末记录当前结果。",
)
set_single_run(
    find_one(document, "下表记录审查阶段实际执行的检查。"),
    "下表记录审查阶段实际执行的检查，保留为修复前基线。"
    "修复后的全量结果见文末“修复与复核记录”。",
)
set_single_run(
    find_one(document, "下列调用链和功能缺口已被定位，"),
    "下列四项补充问题均已完成最小范围修复并添加回归测试。"
    "历史会话修复范围为可发现、可查看和可导出；删除与自动保留仍在产品规划中。",
)
set_single_run(
    find_one(document, "交付内容包括审查结果、复核用例和已完成修复的记录。"),
    "本报告按发现逐项追加“【已修复】”标记，并保留修复前证据。"
    "各项代码修改先添加可复现测试，再执行针对性回归与全项目检查。",
)

document.add_heading("修复与复核记录（截至 2026 年 9 月 5 日）", level=1)
document.add_paragraph(
    "本节记录修复后的源码状态。“【已修复】”表示对应缺陷已有代码变更和自动化回归；"
    "它不表示历史 1.0.4 安装包已经重建，也不替代真实 Windows 安装与升级验收。"
)
document.add_heading("P1 修复记录", level=2)
for text in (
    "R01：结束保护改用安全桌面的 Windows 系统凭据窗口；服务校验当前用户 SID、短期挑战、连接身份、会话和重放边界。取消或失败时继续保护。",
    "R02：重点目录配置接入服务和界面，生产会话启动文件采集器；目录范围限制为最多 16 个本地路径。",
    "R03：采集故障、恢复和健康状态进入服务 API 与界面，关闭阶段先等待采集器收尾，再排空 Writer。",
    "R04：Preparing、Active、Finalizing、Completed 与采集启动、结束扫描和持久化提交对齐，按钮按真实中间状态工作。",
    "R06：服务恢复活动会话时转为 Degraded，并写入 service_recovered_after_interruption 边界事件；该事件计入风险和报告的采集缺口。",
    "R08：ProcessKey、ObjectKey 采用 AES-GCM 元数据封装，旧明文字段迁移到 schema v4，并执行 WAL 截断和空间回收。",
    "R09：数据目录创建前核对整棵目录所有者和重解析点；受控创建后设置 owner、保护 DACL，并在离线维护前复核写权限。",
    "R10：MSI 卸载维护动作以提升身份运行。仅 LocalSystem 可委托安装所有者 SID，普通管理员无法通过参数替代真实所有者校验。",
    "R11：MSI 接入准备、应用、停止回滚、恢复和提交事务；维护程序以 MSI Binary 运行，并覆盖文件、数据、服务配置、策略、清单和启动入口恢复。",
):
    document.add_paragraph(text)

document.add_heading("P2 修复记录", level=2)
for text in (
    "R12：发布、预检、安装、升级与验证统一覆盖 Service、UI、Agent、Maintenance 四组件；清单 v2 兼容读取旧 v1。",
    "R13：维护工具可在服务缺失或停止时读取受保护持久化状态，失败阶段写入结果并保留可重试事务与备份。",
    "R14：界面恢复 Draft、Preparing、Finalizing 等中间状态，继续对应启动或收尾流程，避免重复创建会话。",
    "R15：注册表通知配合每秒校正扫描，覆盖重新订阅和事件写入窗口内的变化。",
    "R16：时间线只返回 16 KiB 负载预览，页面受 512 KiB 编码预算约束；原始负载不随页面返回。",
    "R17：报告按完整消息 JSON 大小选择直接响应或分块下载，覆盖 HTML 实体、反斜杠和换行膨胀。",
    "R18：报告支持事件序号范围，单次限制 100000 条事件和 32 MiB 已选原始负载，界面展示并发送范围。",
    "R19：时间线使用索引化键集分页，每页最多扫描 4096 条，仅对当前页验链；完整报告仍验证整个会话，连接取消会停止后台查询。",
    "R20：首次密钥使用排他发布，同进程和跨实例竞争均返回同一落盘密钥。",
    "R21：已有服务重配时保留带引号的可执行文件路径，避免 Program Files 路径解析歧义。",
    "R22：组件激活与恢复失败保留完整错误链、备份位置和未完成事务，再次维护会识别遗留状态。",
    "R23：MSI 脚本强制接收 40 位证书指纹，隔离暂存并完成四组件签名、finalize、MSI 签名和回执后才发布；失败不会覆盖现有产物。",
    "R24：已应用筛选与正在编辑的筛选分开保存，加载更多始终使用生成游标时的筛选。",
    "R25：风险摘要从完整会话计算事件和采集缺口，不受时间线页数或界面筛选影响。",
    "R26：HTML 与 Markdown 模板在用户开启选项时输出进程关联键，关闭时不输出。",
    "R27：Windows 服务不再声明 Pause 和 Continue，SCM 状态不会与仍在运行的采集器相矛盾。",
):
    document.add_paragraph(text)

document.add_heading("补充问题修复记录", level=2)
for text in (
    "S01：数据目录维护锁覆盖空闲检查、停服、数据备份、替换、健康验证和回滚；持久化维护日志同时阻止新会话。运行中的旧服务缺少冻结能力时安全拒绝维护。",
    "S02：安装、部署、升级和卸载只接受固定服务名及默认 ProgramData 目录，不受支持的自定义值在文件复制或系统修改前返回错误。",
    "S03：新增稳定的历史会话键集分页协议、服务 API、桌面桥接和界面入口。用户可在当前会话与历史会话之间切换并查看、分析和导出。单页同时受 50 条与 512 KiB 编码预算约束。",
    "S04：文件重命名可跨通知批次配对；溢出、启动订阅窗口、采集器重启和结束阶段执行有界元数据扫描，扫描失败写入采集缺口。",
):
    document.add_paragraph(text)

document.add_heading("文档与交付边界", level=2)
document.add_paragraph(
    "README、产品需求文档和技术架构文档已经区分架构目标、当前源码和未交付能力，"
    "并统一四组件、CGO、实际存储路径、进程快照、共享主密钥、报告格式、当前规则重算、"
    "历史入口和正式 WXS 来源等描述。项目 Git 根仍位于上级用户目录，未创建新的仓库边界。"
)
document.add_paragraph(
    "Session Agent 前台窗口与输入活动、ETW 进程流、USN 与文件读取审核、完整系统基线、"
    "历史删除和自动保留、JSON 与独立校验包、每会话独立密钥、固定历史规则结果仍未交付。"
    "这些属于已明确记录的产品范围，未用缺陷标记掩盖。"
)

document.add_heading("修复后验证", level=2)
for text in (
    "Go：go test -race ./... -count=1 通过；go vet ./... 通过。",
    "前端：7 个测试文件中的 30 项测试通过；TypeScript 类型检查和 Vite 生产构建通过。",
    "规模：10 万和百万事件的键集分页测试通过；历史会话测试覆盖 53 条跨页、重启可发现、编码预算和超大单项拒绝。",
    "维护：签名流程故障注入通过；WiX 临时不可安装夹具完成编译、ICE 和安装表检查，未执行产品安装。构建无 WiX 或 ICE 警告，四个未版本化 EXE 进入 MsiFileHash，快捷方式不跨组件引用文件。",
    "发布：现有 1.0.4 的四个 EXE 和 MSI 在本机测试证书链下显示 Valid，签名证书为 Desktop Guard Pro Test Signing，未发现时间戳。本轮未重建这些产物。",
    "隔离机：正式证书签名、普通 UAC 卸载、真实服务控制、升级回滚、文件占用、重启清理、跨 Windows 版本和 Hello 生物识别仍需执行。",
):
    document.add_paragraph(text)

for field in (
    "title", "subject", "author", "last_modified_by", "comments", "keywords",
    "category", "identifier", "language", "content_status", "version",
):
    setattr(document.core_properties, field, "")

buffer = io.BytesIO()
document.save(buffer)
temporary = None
try:
    with tempfile.NamedTemporaryFile(
        dir=REPORT.parent, prefix=".review-complete-", suffix=".docx", delete=False,
    ) as output:
        temporary = Path(output.name)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root:
                    if etree.QName(node).localname in (
                        "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
                    ):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)

    verified = Document(temporary)
    if len(verified.paragraphs) <= original_count or len(verified.tables) != original_tables:
        raise AssertionError("report structure was not extended as expected")
    if [section._sectPr.xml for section in verified.sections] != original_sections:
        raise AssertionError("page layout changed")
    for paragraph in verified.paragraphs[:original_count]:
        old_properties = original_properties.get(paragraph.text)
        if old_properties is not None and (
            paragraph._p.pPr.xml if paragraph._p.pPr is not None else None
        ) != old_properties:
            raise AssertionError(f"paragraph formatting changed: {paragraph.text}")
    headings = [paragraph for paragraph in verified.paragraphs if re.match(r"^R\d{2}  ", paragraph.text)]
    if len(headings) != 25 or any("【已修复】" not in paragraph.text for paragraph in headings):
        raise AssertionError("one or more review findings lack the repaired marker")
    for prefix in SUPPLEMENTAL_PREFIXES:
        if "【已修复】" not in find_one(verified, prefix).text:
            raise AssertionError(f"supplemental finding lacks the repaired marker: {prefix}")
    appended = "\n".join(paragraph.text for paragraph in verified.paragraphs[original_count:])
    if any(f"{finding}：" not in appended for finding in FIXED_IDS):
        raise AssertionError("one or more repair records are missing")
    if not any(paragraph.text == "修复后验证" for paragraph in verified.paragraphs):
        raise AssertionError("verification section is missing")
    if any(
        getattr(verified.core_properties, field)
        for field in ("title", "subject", "author", "last_modified_by", "comments", "keywords", "category")
    ):
        raise AssertionError("descriptive metadata remains")
    with zipfile.ZipFile(temporary) as archive:
        root = etree.fromstring(archive.read("docProps/app.xml"))
        if any(
            node.text for node in root
            if etree.QName(node).localname in (
                "Application", "Company", "Manager", "Template", "AppVersion", "HyperlinkBase",
            )
        ):
            raise AssertionError("application metadata remains")
    os.replace(temporary, REPORT)
    temporary = None
finally:
    if temporary is not None:
        temporary.unlink(missing_ok=True)

print("Review report repair records appended and verified.")
